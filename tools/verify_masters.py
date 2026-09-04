"""Excel round trip on the Master Data page: download the template, upload an edited workbook (1 new row, 1 update,
1 invalid reference, 1 missing key), check validation preview, apply, badges and audit. Headless Chrome."""
import sys, time, subprocess, os, tempfile
from playwright.sync_api import sync_playwright
from openpyxl import Workbook, load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); PORT = 8112
srv = subprocess.Popen([sys.executable, "-m", "http.server", str(PORT)], cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
res = []
def check(n, c, extra=""): res.append(bool(c)); print(("PASS " if c else "FAIL ") + n + (("  " + extra) if extra and not c else ""))
try:
    with sync_playwright() as p:
        b = p.chromium.launch(channel="chrome", headless=True); ctx = b.new_context(accept_downloads=True, viewport={"width": 1600, "height": 1000}); pg = ctx.new_page()
        errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
        for i in range(20):
            try: pg.goto(f"http://localhost:{PORT}/masters.html#workCentres", wait_until="load"); break
            except Exception: time.sleep(0.4)
        pg.wait_for_timeout(800); pg.evaluate("localStorage.removeItem('khp-demo-state')"); pg.reload(); pg.wait_for_timeout(800)
        check("masters page renders work centres", "KHP-CGL-01" in pg.evaluate("document.body.innerText") and pg.evaluate("typeof XLSX") == "object")
        with pg.expect_download() as dl: pg.click("#dl")
        path = os.path.join(tempfile.gettempdir(), "khp_wc_template.xlsx"); dl.value.save_as(path)
        wb = load_workbook(path); ws = wb[wb.sheetnames[0]]; hdr = [c.value for c in ws[1]]
        check("template downloaded with header + README sheet", "id" in hdr and "sapWorkCentre" in hdr and "README" in wb.sheetnames, str(hdr))
        # edit: update CGL capacity, add a new work centre, add a bad reference row, add a row without key
        col = {h: i + 1 for i, h in enumerate(hdr)}
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, col["id"]).value == "KHP-CGL-01": ws.cell(r, col["capacityTph"]).value = 24
        def add(**kw):
            row = [kw.get(h, "") for h in hdr]; ws.append(row)
        add(id="KHP-CGL2-02", lineId="CGL2", name="CGL 2 exit inspection", sapWorkCentre="WC_CGL2B", costCentre="CC-KHP-201", capacityTph=25, shiftsPerDay=3, bypassAllowed="no", status="PLANNED")
        add(id="KHP-XYZ-01", lineId="XYZ", name="bad line", sapWorkCentre="WC_X", capacityTph="abc", status="ACTIVE")
        add(lineId="PKL", name="no key", sapWorkCentre="WC_PKL9")
        edited = os.path.join(tempfile.gettempdir(), "khp_wc_edited.xlsx"); wb.save(edited)
        pg.set_input_files("#file", edited); pg.wait_for_timeout(1200)
        t = pg.evaluate("document.getElementById('pmeta').innerText")
        check("preview: 1 new, 1 update, 2 errors", "1 new" in t and "1 updates" in t and "2 errors" in t, t)
        prev = pg.evaluate("document.getElementById('prev').innerText")
        check("preview: reference + numeric + missing-key errors shown", 'lineId "XYZ" not in lines' in prev and "must be numeric" in prev and "missing id" in prev)
        pg.click("#apply"); pg.wait_for_timeout(600)
        body = pg.evaluate("document.body.innerText")
        check("applied: NEW badge, UPDATED badge, audit row", "KHP-CGL2-02" in body and "NEW" in body and "UPDATED" in body and "khp_wc_edited.xlsx" in body)
        check("state persisted for other pages", pg.evaluate("Object.keys(KHPState.get('masters',{})).indexOf('workCentres')>=0"))
        pg.click("[data-m=compatibility]"); pg.wait_for_timeout(400); check("matrix view for compatibility", "matrix view" in pg.evaluate("document.body.innerText").lower())
        pg.click("[data-m=roles]"); pg.wait_for_timeout(300); check("roles master lists page rights", "Transact" in pg.evaluate("document.body.innerText") or "V/S/T" in pg.evaluate("document.body.innerText"))
        check("zero JS errors", not errs, str(errs[:2]))
        b.close()
    ok = all(res); print("MASTERS", "PASS" if ok else "FAIL", f"({sum(res)}/{len(res)})"); sys.exit(0 if ok else 1)
finally:
    srv.kill()
